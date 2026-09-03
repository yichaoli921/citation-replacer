#!/usr/bin/env python3
"""
replace_citations.py
跨 run 替换 docx 中的 author-year 引用为上标 [N-M] / [N，M，K-L]

依赖：python-docx
运行：python replace_citations.py <input.docx> <cite_map.tsv> [style_name] [output.docx]
"""

import sys
import re
import csv
import yaml
import glob
import shutil
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn


# ============================================================
# 默认风格：《经济地理》
# ============================================================
DEFAULT_STYLE = {
    'name': 'jingji-dili',
    'citation': {
        'bracket_open': '［',
        'bracket_close': '］',
        'superscript': True,
        'range_separator': '-',
        'discontinuous_separator': '，',
    },
    'narrative': {
        'enabled': True,
        'format': '{author}［{ref}］',
        'superscript': False,
    },
}


def load_style(name: str) -> dict:
    """加载期刊风格定义"""
    base = Path(__file__).parent.parent / 'references' / 'journal-styles' / f'{name}.md'
    if not base.exists():
        print(f'⚠️  未找到风格文件 {base}，使用默认《经济地理》风格')
        return DEFAULT_STYLE
    # 解析 markdown 中的 yaml 块
    text = base.read_text()
    m = re.search(r'```yaml\n(.*?)\n```', text, re.DOTALL)
    if not m:
        print(f'⚠️  风格文件无 yaml 块，使用默认')
        return DEFAULT_STYLE
    style = yaml.safe_load(m.group(1))
    style['name'] = name
    return style


# ============================================================
# 引用映射表
# ============================================================
def load_cite_map(tsv_path: str) -> dict:
    """
    读取 cite_map.tsv 格式：author\tyear\tnumber
    返回 {(normalized_author, year): number, ...}
    """
    mapping = {}
    with open(tsv_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f, delimiter='\t')
        header = next(reader, None)
        for row in reader:
            if len(row) < 3:
                continue
            author, year, num = row[0].strip(), int(row[1].strip()), int(row[2].strip())
            mapping[normalize_author(author), year] = num
    return mapping


def normalize_author(name: str) -> str:
    """作者名归一化"""
    name = re.sub(r'\s+', '', name)
    name = name.replace('．', '.').replace('·', '')
    return name.lower()


# ============================================================
# 引用块识别与解析
# ============================================================
CITE_PAR = re.compile(
    r'[（(]'
    r'(?=[^）)]*[一-龥A-Za-z])'
    r'[^）)]*?'
    r'(?<!\d)(?:19|20)\d{2}(?![0-9])'
    r'[）)]'
)

NARRATIVE_PAT = re.compile(
    r'(?<![0-9．。])'
    r'([一-龥A-Za-z][一-龥 A-Za-z]{1,30}?(?:等|et\s+al\.?)?)'
    r'[（(]'
    r'(\d{4}[a-z]?(?:[；;]\s*\d{4}[a-z]?)*)'
    r'[）)]'
)


def is_citation_block(text: str) -> bool:
    """判断是否包含真正的引用"""
    has_author = re.search(r'[一-龥A-Za-z]', text) is not None
    has_year = re.search(r'\b(?:19|20)\d{2}\b', text) is not None
    if not (has_author and has_year):
        return False
    excluded = ['中国', '邮编', '通讯作者', '邮箱', '@', 'http', 'www.', 'GS(', '数据来源']
    if any(s in text for s in excluded):
        return False
    return True


def parse_single_cite(s: str):
    """'作者，YYYY' → (author, year)"""
    # OCR 容错：去掉"鉴"等误字前缀（用 lookahead 检测后面是中文字符）
    s = re.sub(r'^[鉴认议谨让论误错谬驾]+(?=[一-龥A-Za-z])', '', s).strip()
    m = re.match(r'^\s*(.+?)\s*[,，]?\s*((?:19|20)\d{2})[a-z]?\s*$', s)
    if m:
        return m.group(1).strip(), int(m.group(2))
    return None


def parse_cite_block(text: str):
    """'A，2020; B and C，2021' → [(author, year), ...]
    注意：不能简单用 '，' 拆分（会把年份切掉）。先按 ;； 拆，每块作为一个"作者-年份"对。
    """
    parts = re.split(r'[;；]', text)
    results = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        # 整个 part 应该是 "作者A和作者B，YYYY"
        # 逗号可选（OCR 容错）
        m = re.match(r'^(.+?)[,，]?\s*((?:19|20)\d{2})[a-z]?\s*$', p)
        if not m:
            continue
        author_part = m.group(1).strip()
        year = int(m.group(2))
        results.append((author_part, year))
    return results


def lookup_number(author: str, year: int, cite_map: dict) -> int | None:
    """查表"""
    norm = normalize_author(author)
    return cite_map.get((norm, year))


def render_ref(numbers: list[int], style: dict) -> str:
    """[1,2,3,5] → '[1-3，5]' 或 '[1-3, 5]'"""
    open_b = style['citation']['bracket_open']
    close_b = style['citation']['bracket_close']
    range_sep = style['citation']['range_separator']
    disc_sep = style['citation']['discontinuous_separator']
    nums = sorted(set(n for n in numbers if n))
    if not nums:
        return ''
    parts = []
    i = 0
    while i < len(nums):
        start = nums[i]
        end = start
        while i + 1 < len(nums) and nums[i+1] == end + 1:
            i += 1
            end = nums[i]
        if start == end:
            parts.append(str(start))
        else:
            parts.append(f'{start}{range_sep}{end}')
        i += 1
    return f'{open_b}{disc_sep.join(parts)}{close_b}'


# ============================================================
# 跨 run 文本拼接
# ============================================================
def get_paragraph_full_text(paragraph):
    """返回完整文本 + 每字符对应 run 索引"""
    runs = list(paragraph.runs)
    text = ''
    char_run_idx = []
    for i, run in enumerate(runs):
        for _ in run.text:
            char_run_idx.append(i)
        text += run.text
    return runs, text, char_run_idx


def apply_replacement(paragraph, runs, start: int, end: int, ref_str: str,
                      superscript: bool = True):
    """在 [start, end) 范围内替换为 ref_str（新建上标 run）"""
    char_run_idx = []
    for i, run in enumerate(runs):
        for _ in run.text:
            char_run_idx.append(i)

    affected = sorted(set(char_run_idx[start:end]))
    if not affected:
        return False
    first_run = runs[affected[0]]
    last_run = runs[affected[-1]]

    # 计算在 first_run/last_run 内的偏移
    start_in_first = start - sum(len(runs[i].text) for i in range(affected[0]))
    end_in_last = end - sum(len(runs[i].text) for i in range(affected[-1]))

    # 创建新 run（先 add 到末尾）
    new_run = paragraph.add_run(ref_str)
    # 复制 first_run 的 rPr（保留字号字体）
    rPr_src = first_run._element.find(qn('w:rPr'))
    if rPr_src is not None:
        new_run._element.insert(0, deepcopy(rPr_src))
    if superscript:
        new_run.font.superscript = True
    # 移动到正确位置
    first_run._element.addnext(new_run._element)

    # 修改受影响的 run 文本
    if first_run is last_run:
        t = first_run.text
        first_run.text = t[:start_in_first] + t[end_in_last:]
    else:
        first_run.text = first_run.text[:start_in_first]
        for ri in affected[1:-1]:
            runs[ri].text = ''
        last_run.text = last_run.text[end_in_last:]
    return True


# ============================================================
# 主处理流程
# ============================================================
def process_paragraph(paragraph, cite_map: dict, style: dict, mode: str = 'both'):
    """
    处理单个段落
    mode: 'parenthetical' / 'narrative' / 'both'
    """
    runs, full_text, _ = get_paragraph_full_text(paragraph)
    if not is_citation_block(full_text):
        return 0

    replacements = []

    # 括号式引用
    if mode in ('parenthetical', 'both'):
        for m in CITE_PAR.finditer(full_text):
            inner = m.group(0)[1:-1]
            parsed = parse_cite_block(inner)
            if not parsed:
                continue
            nums = [lookup_number(a, y, cite_map) for a, y in parsed]
            if any(n is None for n in nums):
                continue
            ref_str = render_ref([n for n in nums if n], style)
            if not ref_str:
                continue
            replacements.append((m.start(), m.end(), ref_str, True))  # superscript=True

    # 叙述式引用
    if mode in ('narrative', 'both') and style.get('narrative', {}).get('enabled'):
        for m in NARRATIVE_PAT.finditer(full_text):
            author = m.group(1).strip()
            years_str = m.group(2)
            # 多个年份
            for ym in re.finditer(r'\d{4}', years_str):
                year = int(ym.group(0))
                num = lookup_number(author, year, cite_map)
                if num is None:
                    continue
                # 替换 (YYYY) 部分
                y_start = m.start(2) + ym.start()
                y_end = m.start(2) + ym.end()
                # 包含括号
                paren_start = y_start - 1
                paren_end = y_end + 1
                if paren_start < 0 or paren_end > len(full_text):
                    continue
                if full_text[paren_start] not in '（(' or full_text[paren_end - 1] not in '）)':
                    # 实际上是括号外的范围，要 trim
                    paren_start = y_start
                    paren_end = y_end
                # 重新定位括号
                # 简化：直接定位 (YYYY) 字符串
                ytext = f'{full_text[y_start-1]}{years_str}{full_text[y_end]}'
                # 找真实括号边界
                # (YYYY)
                abs_start = full_text.rfind('（', m.start(), y_start + 1)
                if abs_start < 0:
                    abs_start = full_text.rfind('(', m.start(), y_start + 1)
                abs_end = full_text.find('）', y_end, m.end())
                if abs_end < 0:
                    abs_end = full_text.find(')', y_end, m.end())
                if abs_start < 0 or abs_end < 0:
                    continue
                ref_str = render_ref([num], style)
                replacements.append((abs_start, abs_end + 1, ref_str, False))  # superscript=False

    if not replacements:
        return 0

    # 倒序处理避免索引偏移
    count = 0
    # 去重：同一范围不要处理两次
    replacements = list({(s, e, r, sup) for s, e, r, sup in replacements})
    replacements.sort(key=lambda x: x[0])
    for start, end, ref_str, sup in reversed(replacements):
        # 注意：随着前面替换，runs 列表会变（new_run 加到末尾），但 char_run_idx 用 runs 拍快照时的索引
        if apply_replacement(paragraph, runs, start, end, ref_str, sup):
            count += 1
    return count


def process_document(doc: Document, cite_map: dict, style: dict,
                     stop_at: str = '参考文献：', mode: str = 'both') -> int:
    """处理整个文档，到 stop_at 段落停止"""
    total = 0
    for p in doc.paragraphs:
        if p.text.strip() == stop_at:
            break
        total += process_paragraph(p, cite_map, style, mode=mode)
    return total


# ============================================================
# 入口
# ============================================================
def main():
    if len(sys.argv) < 3:
        print('用法：python replace_citations.py <input.docx> <cite_map.tsv> [style] [output.docx]')
        print('示例：python replace_citations.py paper.docx cite_map.tsv jingji-dili paper_replaced.docx')
        sys.exit(1)

    input_docx = sys.argv[1]
    cite_tsv = sys.argv[2]
    style_name = sys.argv[3] if len(sys.argv) > 3 else 'jingji-dili'
    output_docx = sys.argv[4] if len(sys.argv) > 4 else input_docx.replace('.docx', '_replaced.docx')

    # 支持 glob
    if '*' in input_docx:
        matches = glob.glob(input_docx)
        if not matches:
            print(f'❌ 未找到匹配 {input_docx} 的文件')
            sys.exit(1)
        input_docx = matches[0]
        print(f'→ 匹配文件：{input_docx}')

    style = load_style(style_name)
    cite_map = load_cite_map(cite_tsv)
    print(f'→ 风格：{style_name}，映射表 {len(cite_map)} 条')

    # 备份
    bak_path = input_docx + '.bak'
    if not Path(bak_path).exists():
        shutil.copy(input_docx, bak_path)
        print(f'→ 已备份到 {bak_path}')

    doc = Document(input_docx)
    count = process_document(doc, cite_map, style)
    print(f'✅ 完成 {count} 处替换')

    doc.save(output_docx)
    print(f'→ 输出：{output_docx}')


if __name__ == '__main__':
    main()
